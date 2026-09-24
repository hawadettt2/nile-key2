from fastapi import APIRouter, Depends, HTTPException, Query
from pydantic import BaseModel
from typing import Optional, List
import os
from dotenv import load_dotenv
import io
import requests
from openpyxl import load_workbook

from app.routers.auth import get_current_user

load_dotenv(dotenv_path=os.path.join(os.path.dirname(__file__), '..', '..', '..', '.env'))

router = APIRouter(prefix="/api/v1/potential-customers", tags=["Potential Customers"])

GOOGLE_DRIVE_API_KEY = os.environ.get("GOOGLE_DRIVE_API_KEY")
ROOT_FOLDER_ID = "1x0yx_x9QYtSNim2a7Loy4mjrHCu7Z8zQ"
GOOGLE_DRIVE_API_BASE = "https://www.googleapis.com/drive/v3"


class ExcelFileResponse(BaseModel):
    id: str
    name: str
    source: Optional[str] = None
    sheet_count: Optional[int] = None
    row_count: Optional[int] = None
    size_bytes: Optional[int] = None
    download_url: Optional[str] = None
    view_url: Optional[str] = None
    created_time: Optional[str] = None
    modified_time: Optional[str] = None


class CountryFolderResponse(BaseModel):
    id: str
    name: str
    file_count: int
    files: List[ExcelFileResponse]


class PotentialCustomersIndexResponse(BaseModel):
    countries: List[CountryFolderResponse]
    total_countries: int
    total_files: int
    source: str
    error: Optional[str] = None


class SheetMeta(BaseModel):
    name: str
    row_count: int
    column_count: int


class ExcelContentResponse(BaseModel):
    file_id: str
    file_name: str
    sheets: List[SheetMeta]
    active_sheet: Optional[str] = None
    rows: List[dict]


def _drive_get(path: str, params: Optional[dict] = None) -> dict:
    if not GOOGLE_DRIVE_API_KEY:
        raise RuntimeError("GOOGLE_DRIVE_API_KEY is not configured on the server.")
    url = f"{GOOGLE_DRIVE_API_BASE}/{path}"
    params = dict(params or {})
    params["key"] = GOOGLE_DRIVE_API_KEY
    params["fields"] = "nextPageToken, files(id, name, mimeType, size, createdTime, modifiedTime, webViewLink, webContentLink)"
    response = requests.get(url, params=params, timeout=30)
    if response.status_code != 200:
        raise RuntimeError(f"Google Drive API error: {response.status_code} - {response.text}")
    return response.json()


def _list_countries() -> List[CountryFolderResponse]:
    data = _drive_get(
        "files",
        params={
            "q": f"'{ROOT_FOLDER_ID}' in parents and mimeType='application/vnd.google-apps.folder' and trashed=false",
            "pageSize": 100,
        },
    )
    countries = []
    for folder in data.get("files", []):
        countries.append(
            CountryFolderResponse(
                id=folder["id"],
                name=folder["name"],
                file_count=0,
                files=[],
            )
        )
    return countries


def _list_country_files(country_id: str) -> List[ExcelFileResponse]:
    data = _drive_get(
        "files",
        params={
            "q": f"'{country_id}' in parents and trashed=false and (mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' or mimeType='text/plain')",
            "pageSize": 100,
        },
    )
    files = []
    for f in data.get("files", []):
        size = f.get("size")
        files.append(
            ExcelFileResponse(
                id=f["id"],
                name=f["name"],
                source="Google Drive",
                size_bytes=int(size) if size else None,
                download_url=f.get("webContentLink") or f"https://drive.google.com/uc?id={f['id']}",
                view_url=f.get("webViewLink") or f"https://drive.google.com/file/d/{f['id']}/view",
                created_time=f.get("createdTime"),
                modified_time=f.get("modifiedTime"),
            )
        )
    return files


def _load_index_from_drive() -> PotentialCustomersIndexResponse:
    countries = _list_countries()
    total_files = 0
    for country in countries:
        files = _list_country_files(country.id)
        country.file_count = len(files)
        country.files = files
        total_files += len(files)
    return PotentialCustomersIndexResponse(
        countries=countries,
        total_countries=len(countries),
        total_files=total_files,
        source="google-drive",
    )


def _fetch_excel_from_drive(file_id: str) -> tuple[str, bytes]:
    if not GOOGLE_DRIVE_API_KEY:
        raise RuntimeError("GOOGLE_DRIVE_API_KEY is not configured on the server.")
    url = f"{GOOGLE_DRIVE_API_BASE}/files/{file_id}"
    params = {"key": GOOGLE_DRIVE_API_KEY, "alt": "media"}
    response = requests.get(url, params=params, timeout=60)
    if response.status_code != 200:
        raise RuntimeError(f"Google Drive API error: {response.status_code} - {response.text}")
    content_type = response.headers.get("Content-Type", "application/octet-stream")
    return content_type, response.content


@router.get("/countries", response_model=PotentialCustomersIndexResponse)
def list_countries(current_user: dict = Depends(get_current_user)):
    try:
        return _load_index_from_drive()
    except RuntimeError as exc:
        error_message = str(exc)
        if "GOOGLE_DRIVE_API_KEY" in error_message:
            raise HTTPException(status_code=503, detail="Google Drive API key is not configured on the server.")
        raise HTTPException(status_code=502, detail=error_message)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to load potential customers from Google Drive: {exc}")


@router.get("/countries/{country_id}", response_model=CountryFolderResponse)
def get_country(country_id: str, current_user: dict = Depends(get_current_user)):
    try:
        index = _load_index_from_drive()
    except RuntimeError as exc:
        error_message = str(exc)
        if "GOOGLE_DRIVE_API_KEY" in error_message:
            raise HTTPException(status_code=503, detail="Google Drive API key is not configured on the server.")
        raise HTTPException(status_code=502, detail=error_message)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to load potential customers from Google Drive: {exc}")

    for country in index.countries:
        if country.id == country_id:
            return country
    raise HTTPException(status_code=404, detail="Country not found")


@router.get("/files/{file_id}/content", response_model=ExcelContentResponse)
def get_excel_content(file_id: str, sheet: Optional[str] = Query(default=None), current_user: dict = Depends(get_current_user)):
    try:
        content_type, content = _fetch_excel_from_drive(file_id)
    except RuntimeError as exc:
        error_message = str(exc)
        if "GOOGLE_DRIVE_API_KEY" in error_message:
            raise HTTPException(status_code=503, detail="Google Drive API key is not configured on the server.")
        raise HTTPException(status_code=502, detail=error_message)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Failed to fetch file from Google Drive: {exc}")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Failed to parse Excel file: {exc}")

    sheet_names = workbook.sheetnames
    active_sheet = sheet or sheet_names[0] if sheet_names else None
    if active_sheet not in sheet_names:
        raise HTTPException(status_code=400, detail=f"Sheet '{active_sheet}' not found. Available sheets: {sheet_names}")

    ws = workbook[active_sheet]
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append({str(idx + 1): (cell if cell is not None else "") for idx, cell in enumerate(row)})

    workbook.close()

    return ExcelContentResponse(
        file_id=file_id,
        file_name="",
        sheets=[
            SheetMeta(name=name, row_count=0, column_count=0)
            for name in sheet_names
        ],
        active_sheet=active_sheet,
        rows=rows,
    )
