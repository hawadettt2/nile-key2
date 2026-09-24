import csv
import io
import json
import uuid
from datetime import datetime, timedelta
from typing import Optional, Dict, Any, List

from app.schemas.customer import (
    CustomerCreate,
    CustomerUpdate,
    CustomerProductResponse,
    CustomerEvidenceResponse,
    CustomerSourceBatchResponse,
    CustomerRawRecordResponse,
    ImportPreviewResponse,
    ImportConfirmRequest,
    ImportConfirmResponse,
    ImportResponse,
    CustomerDetail,
)
from app.services.audit import log_audit
from app.schemas.audit import AuditLogCreate
from app.core.database import get_db, DatabaseSession
from app.services.base import now_iso, parse_json, dumps_json


def _now_iso() -> str:
    return now_iso()


def _customer_row_to_response(row: dict) -> dict:
    response = {}
    for key in [
        "id", "email", "phone", "mobile", "whatsapp", "website",
        "address", "city", "country", "tax_id", "import_license", "category",
        "notes", "status", "data_status", "verification_status", "crm_status",
        "activity_status", "activity_window_label", "created_at", "updated_at", "created_by"
    ]:
        response[key] = row.get(key)
    response["name"] = row.get("name") if row.get("name") is not None else row.get("company_name")
    response["contact_person"] = row.get("contact_person") if row.get("contact_person") is not None else row.get("contact_name")
    response["name_en"] = row.get("name_en")
    response["activity_window_start"] = row.get("activity_window_start")
    response["activity_window_end"] = row.get("activity_window_end")
    return response


def _build_customer_detail(row: dict, products: list, evidence: list, batches: list, raw_records: list) -> dict:
    base = _customer_row_to_response(row)
    base["products"] = products
    base["evidence"] = evidence
    base["source_batches"] = batches
    base["raw_records"] = raw_records
    return base


def list_customers(
    search: Optional[str] = None,
    status: Optional[str] = None,
    country: Optional[str] = None,
    category: Optional[str] = None,
    data_status: Optional[str] = None,
    verification_status: Optional[str] = None,
    crm_status: Optional[str] = None,
    activity_status: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
) -> list[dict]:
    from app.services.base import build_list_query

    conn = get_db()
    try:
        filters = {
            "status": status,
            "country": country,
            "category": category,
            "data_status": data_status,
            "verification_status": verification_status,
            "crm_status": crm_status,
            "activity_status": activity_status,
        }
        query, params = build_list_query(
            "customers",
            filters={k: v for k, v in filters.items() if v is not None},
            search_fields=["name", "name_en", "email", "phone", "contact_person"],
            search=search,
            order_by="created_at DESC",
            limit=limit,
            offset=skip,
        )
        session = DatabaseSession(conn)
        rows = session.fetch_all(query, tuple(params))
        return [_customer_row_to_response(dict(r)) for r in rows]
    finally:
        conn.close()


def get_customer(customer_id: int) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        row = session.fetch_one("SELECT * FROM customers WHERE id = ?", (customer_id,))
        if not row:
            raise ValueError("Customer not found")
        products = session.fetch_all(
            "SELECT id, product_description, hs_code, hs_code_description, quantity, unit, created_at FROM customer_products WHERE customer_id = ? ORDER BY id ASC",
            (customer_id,),
        )
        evidence = session.fetch_all(
            "SELECT id, evidence_type, evidence_data, observed_at, activity_window_start, activity_window_end, activity_window_label, created_at FROM customer_evidence WHERE customer_id = ? ORDER BY created_at DESC",
            (customer_id,),
        )
        raw_records = session.fetch_all(
            """
            SELECT cr.id, cr.batch_id, cr.sheet_name, cr.row_number, cr.raw_data,
                   cr.validation_errors, cr.conflict_resolution, cr.conflict_details, cr.created_at
            FROM customer_raw_records cr
            WHERE cr.normalized_customer_id = ?
            ORDER BY cr.id ASC
            """,
            (customer_id,),
        )
        batch_ids = [r["batch_id"] for r in raw_records]
        batches = []
        if batch_ids:
            placeholders = ",".join(["?"] * len(batch_ids))
            batches = session.fetch_all(
                f"SELECT id, source_type, source_format, source_name, source_reference, source_url, file_name, row_count, success_count, skip_count, error_count, status, imported_at, imported_by, created_at FROM customer_source_batches WHERE id IN ({placeholders})",
                tuple(batch_ids),
            )
        return _build_customer_detail(
            dict(row),
            [dict(p) for p in products],
            [dict(e) for e in evidence],
            [dict(b) for b in batches],
            [
                {
                    "id": r["id"],
                    "batch_id": r["batch_id"],
                    "sheet_name": r["sheet_name"],
                    "row_number": r["row_number"],
                    "raw_data": parse_json(r["raw_data"]),
                    "validation_errors": parse_json(r["validation_errors"], []),
                    "conflict_resolution": r["conflict_resolution"],
                    "conflict_details": parse_json(r["conflict_details"], {}),
                    "created_at": r["created_at"],
                }
                for r in raw_records
            ],
        )
    finally:
        conn.close()


def create_customer(data: CustomerCreate, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        with session.transaction():
            customer_id = session.insert(
                "customers",
                {
                    "name": data.name,
                    "name_en": data.name_en,
                    "contact_person": data.contact_person,
                    "email": data.email,
                    "phone": data.phone,
                    "mobile": data.mobile,
                    "whatsapp": data.whatsapp,
                    "website": data.website,
                    "address": data.address,
                    "city": data.city,
                    "country": data.country,
                    "tax_id": data.tax_id,
                    "import_license": data.import_license,
                    "category": data.category,
                    "notes": data.notes,
                    "status": "active",
                    "data_status": "normalized",
                    "verification_status": "unverified",
                    "crm_status": "prospect",
                    "activity_status": "unknown",
                    "created_at": now_iso(),
                    "created_by": current_user["id"],
                },
            )
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="create", entity_type="customer", entity_id=customer_id, details=data.name),
        )
        return {"id": customer_id, "message": "Customer created successfully"}
    finally:
        conn.close()


def update_customer(customer_id: int, data: CustomerUpdate, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        existing = session.fetch_one("SELECT id FROM customers WHERE id = ?", (customer_id,))
        if not existing:
            raise ValueError("Customer not found")

        updates = {}
        for field, value in data.model_dump(exclude_unset=True).items():
            if value is not None:
                updates[field] = value
        if not updates:
            return {"message": "No changes"}

        with session.transaction():
            session.update("customers", customer_id, updates)
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="update", entity_type="customer", entity_id=customer_id),
        )
        return {"message": "Customer updated successfully"}
    finally:
        conn.close()


def delete_customer(customer_id: int, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        with session.transaction():
            updated = session.update("customers", customer_id, {"status": "inactive"})
        if not updated:
            return {"message": "No changes"}
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="delete", entity_type="customer", entity_id=customer_id),
        )
        return {"message": "Customer deactivated successfully"}
    finally:
        conn.close()


def get_countries() -> list[str]:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        rows = session.fetch_all("SELECT DISTINCT country FROM customers WHERE country IS NOT NULL ORDER BY country ASC", ())
        return [r["country"] for r in rows if r["country"]]
    finally:
        conn.close()


def list_source_batches(skip: int = 0, limit: int = 100) -> list[dict]:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        rows = session.fetch_all(
            "SELECT * FROM customer_source_batches ORDER BY created_at DESC LIMIT ? OFFSET ?",
            (limit, skip),
        )
        return [dict(r) for r in rows]
    finally:
        conn.close()


def list_products(customer_id: int) -> list[dict]:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        rows = session.fetch_all(
            "SELECT id, product_description, hs_code, hs_code_description, quantity, unit, created_at FROM customer_products WHERE customer_id = ? ORDER BY id ASC",
            (customer_id,),
        )
        return [dict(r) for r in rows]
    finally:
        conn.close()


def create_product(customer_id: int, data: dict, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        with session.transaction():
            product_id = session.insert(
                "customer_products",
                {
                    "customer_id": customer_id,
                    "product_description": data.get("product_description"),
                    "hs_code": data.get("hs_code"),
                    "hs_code_description": data.get("hs_code_description"),
                    "quantity": data.get("quantity"),
                    "unit": data.get("unit"),
                },
            )
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="create", entity_type="customer_product", entity_id=product_id),
        )
        return {"id": product_id, "message": "Product added successfully"}
    finally:
        conn.close()


def delete_product(customer_id: int, product_id: int, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        with session.transaction():
            deleted = session.delete("customer_products", product_id)
        if not deleted:
            return {"message": "No changes"}
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="delete", entity_type="customer_product", entity_id=product_id),
        )
        return {"message": "Product removed successfully"}
    finally:
        conn.close()


def list_evidence(customer_id: int) -> list[dict]:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        rows = session.fetch_all(
            "SELECT id, evidence_type, evidence_data, observed_at, activity_window_start, activity_window_end, activity_window_label, created_at FROM customer_evidence WHERE customer_id = ? ORDER BY created_at DESC",
            (customer_id,),
        )
        return [
            {
                "id": r["id"],
                "evidence_type": r["evidence_type"],
                "evidence_data": parse_json(r["evidence_data"], {}),
                "observed_at": r["observed_at"],
                "activity_window_start": r["activity_window_start"],
                "activity_window_end": r["activity_window_end"],
                "activity_window_label": r["activity_window_label"],
                "created_at": r["created_at"],
            }
            for r in rows
        ]
    finally:
        conn.close()


def create_evidence(customer_id: int, data: dict, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        with session.transaction():
            evidence_id = session.insert(
                "customer_evidence",
                {
                    "customer_id": customer_id,
                    "evidence_type": data.get("evidence_type", "source_specific"),
                    "evidence_data": dumps_json(data.get("evidence_data")),
                    "observed_at": data.get("observed_at"),
                    "activity_window_start": data.get("activity_window_start"),
                    "activity_window_end": data.get("activity_window_end"),
                    "activity_window_label": data.get("activity_window_label"),
                },
            )
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="create", entity_type="customer_evidence", entity_id=evidence_id),
        )
        return {"id": evidence_id, "message": "Evidence added successfully"}
    finally:
        conn.close()


def _parse_csv(content: bytes, header_row: int = 1, data_start: int = 0, data_end: Optional[int] = None) -> tuple[list[dict], list[str]]:
    text = content.decode("utf-8-sig")
    lines = text.splitlines()
    if not lines:
        raise ValueError("empty_csv")

    # header_row is 1-based for user convenience
    header_idx = max(0, header_row - 1)
    if header_idx >= len(lines):
        raise ValueError("invalid_header_row")

    headers = [h.strip() for h in lines[header_idx].split(",")]
    start = max(header_idx + 1, data_start)
    end = data_end if data_end is not None else len(lines)
    rows = []
    for line_idx in range(start, min(end, len(lines))):
        line = lines[line_idx].strip()
        if not line:
            continue
        values = [v.strip() for v in line.split(",")]
        row = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))}
        rows.append(row)
    if not rows:
        raise ValueError("empty_csv")
    return rows, headers


def _parse_xlsx(
    content: bytes,
    header_row: int = 1,
    data_start: int = 0,
    data_end: Optional[int] = None,
) -> tuple[dict[str, list[dict]], list[str]]:
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("xlsx_unsupported") from exc

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets: dict[str, list[dict]] = {}
    all_headers: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[dict] = []
        headers: list[str] = []
        header_idx = max(0, header_row - 1)
        start = max(header_idx + 1, data_start)
        end = data_end if data_end is not None else None
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = list(row)
            if not any(v is not None for v in values):
                continue
            if row_idx == header_row:
                headers = [str(v) if v is not None else f"column_{i+1}" for i, v in enumerate(values)]
                continue
            if row_idx < start:
                continue
            if end is not None and row_idx > end:
                continue
            rows.append({headers[i]: values[i] if i < len(values) else "" for i in range(len(headers))})
        if headers:
            sheets[sheet_name] = rows
            all_headers.extend(headers)
    wb.close()
    if not sheets:
        raise ValueError("empty_xlsx")
    return sheets, list(dict.fromkeys(all_headers))


def _detect_duplicate_candidates(session: DatabaseSession, row: dict, batch_id: int) -> list[dict]:
    candidates: list[dict] = []
    tax_id = (row.get("tax_id") or row.get("import_license") or "").strip()
    email = (row.get("email") or "").strip()
    phone = (row.get("phone") or row.get("mobile") or "").strip()
    website = (row.get("website") or "").strip()
    name = (row.get("name") or row.get("company_name") or "").strip()
    country = (row.get("country") or "").strip()

    if tax_id:
        existing = session.fetch_all(
            "SELECT * FROM customers WHERE LOWER(COALESCE(tax_id, import_license)) = LOWER(?) LIMIT 1",
            (tax_id,),
        )
        if existing:
            candidates.append({"customer": dict(existing[0]), "match_type": "tax_id_or_license"})
    if email:
        existing = session.fetch_all("SELECT * FROM customers WHERE email = ? LIMIT 1", (email,))
        if existing:
            candidates.append({"customer": dict(existing[0]), "match_type": "email"})
    if phone:
        existing = session.fetch_all(
            "SELECT * FROM customers WHERE phone = ? OR mobile = ? LIMIT 1",
            (phone, phone),
        )
        if existing:
            candidates.append({"customer": dict(existing[0]), "match_type": "phone"})
    if website:
        existing = session.fetch_all("SELECT * FROM customers WHERE website = ? LIMIT 1", (website,))
        if existing:
            candidates.append({"customer": dict(existing[0]), "match_type": "website"})
    if name and country:
        existing = session.fetch_all(
            "SELECT * FROM customers WHERE LOWER(name) = LOWER(?) AND LOWER(country) = LOWER(?) LIMIT 1",
            (name, country),
        )
        if existing:
            candidates.append({"customer": dict(existing[0]), "match_type": "name_country"})

    return candidates


def _normalize_customer_from_row(row: dict, source_batch_id: int) -> dict:
    name = row.get("name") or row.get("company_name") or ""
    country = row.get("country") or ""
    if not name or not country:
        raise ValueError("missing_required")

    return {
        "name": name.strip(),
        "name_en": row.get("name_en"),
        "contact_person": row.get("contact_person") or row.get("contact_name"),
        "email": row.get("email"),
        "phone": row.get("phone"),
        "mobile": row.get("mobile"),
        "whatsapp": row.get("whatsapp"),
        "website": row.get("website"),
        "address": row.get("address"),
        "city": row.get("city"),
        "country": country.strip(),
        "tax_id": row.get("tax_id"),
        "import_license": row.get("import_license"),
        "category": row.get("category"),
        "notes": row.get("notes"),
        "status": "active",
        "data_status": "normalized",
        "verification_status": "unverified",
        "crm_status": "prospect",
        "activity_status": "unknown",
        "created_at": now_iso(),
        "created_by": None,
        "raw_data": json.dumps(row, ensure_ascii=False, default=str),
    }


def _create_customer_from_normalized(session: DatabaseSession, normalized: dict, current_user: dict) -> int:
    data = dict(normalized)
    data.pop("raw_data", None)
    data["created_by"] = current_user.get("id")
    customer_id = session.insert("customers", data)
    return customer_id


def import_preview(file: io.BytesIO, filename: str, current_user: dict) -> dict:
    if not filename:
        raise ValueError("filename_required")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext not in {"csv", "xlsx"}:
        raise ValueError("unsupported_format")

    content = file.read()
    if not content:
        raise ValueError("empty_file")

    conn = get_db()
    try:
        session = DatabaseSession(conn)
        now = now_iso()
        expires_at = (datetime.utcnow() + timedelta(hours=24)).isoformat()
        with session.transaction():
            batch_id = session.insert(
                "customer_source_batches",
                {
                    "source_type": "file",
                    "source_format": ext,
                    "source_name": filename,
                    "file_name": filename,
                    "status": "preview",
                    "row_count": 0,
                    "success_count": 0,
                    "skip_count": 0,
                    "error_count": 0,
                    "imported_at": now,
                    "imported_by": current_user.get("id"),
                    "expires_at": expires_at,
                    "mapping_metadata": None,
                },
            )

        # Parse and persist raw records for ALL rows across ALL sheets
        preview_rows = []
        detected_columns = []
        raw_records_to_insert = []

        if ext == "csv":
            rows, headers = _parse_csv(content)
            detected_columns = headers
            for idx, row in enumerate(rows, start=1):
                raw_data = json.dumps(row, ensure_ascii=False, default=str)
                raw_records_to_insert.append({
                    "batch_id": batch_id,
                    "sheet_name": None,
                    "row_number": idx,
                    "raw_data": raw_data,
                    "validation_errors": None,
                    "conflict_resolution": None,
                    "conflict_details": None,
                })
                if idx <= 10:
                    preview_rows.append({"row_number": idx, "raw": row})
        else:
            sheets, headers = _parse_xlsx(content)
            detected_columns = headers
            for sheet_name, rows in sheets.items():
                for idx, row in enumerate(rows, start=1):
                    raw_data = json.dumps(row, ensure_ascii=False, default=str)
                    raw_records_to_insert.append({
                        "batch_id": batch_id,
                        "sheet_name": sheet_name,
                        "row_number": idx,
                        "raw_data": raw_data,
                        "validation_errors": None,
                        "conflict_resolution": None,
                        "conflict_details": None,
                    })
                    if len([r for r in preview_rows if r.get("sheet_name") == sheet_name]) < 10:
                        preview_rows.append({"row_number": idx, "sheet_name": sheet_name, "raw": row})

        if raw_records_to_insert:
            with session.transaction():
                for record in raw_records_to_insert:
                    session.insert("customer_raw_records", record)
                session.update("customer_source_batches", batch_id, {"row_count": len(raw_records_to_insert)})

        return {
            "batch_id": batch_id,
            "file_name": filename,
            "sheet_name": preview_rows[0].get("sheet_name") if preview_rows else None,
            "total_rows": len(preview_rows),
            "preview_rows": preview_rows,
            "detected_columns": detected_columns,
            "target_fields": [
                {"key": "name", "label_en": "Name", "label_ar": "الاسم", "required": True},
                {"key": "country", "label_en": "Country", "label_ar": "الدولة", "required": True},
                {"key": "contact_person", "label_en": "Contact Person", "label_ar": "شخص الاتصال", "required": False},
                {"key": "email", "label_en": "Email", "label_ar": "البريد الإلكتروني", "required": False},
                {"key": "phone", "label_en": "Phone", "label_ar": "الهاتف", "required": False},
                {"key": "mobile", "label_en": "Mobile", "label_ar": "الجوال", "required": False},
                {"key": "whatsapp", "label_en": "WhatsApp", "label_ar": "واتساب", "required": False},
                {"key": "website", "label_en": "Website", "label_ar": "الموقع الإلكتروني", "required": False},
                {"key": "address", "label_en": "Address", "label_ar": "العنوان", "required": False},
                {"key": "city", "label_en": "City", "label_ar": "المدينة", "required": False},
                {"key": "tax_id", "label_en": "Tax ID", "label_ar": "الرقم الضريبي", "required": False},
                {"key": "import_license", "label_en": "Import License", "label_ar": "رخصة الاستيراد", "required": False},
                {"key": "category", "label_en": "Category", "label_ar": "الفئة", "required": False},
                {"key": "product_description", "label_en": "Product Description", "label_ar": "وصف المنتج", "required": False},
                {"key": "hs_code", "label_en": "HS Code", "label_ar": "رمز HS", "required": False},
                {"key": "hs_code_description", "label_en": "HS Description", "label_ar": "وصف HS", "required": False},
            ],
        }
    finally:
        conn.close()


def _apply_field_mapping(raw_data: dict, field_mapping: Dict[str, str] | None) -> dict:
    if not field_mapping:
        return raw_data
    mapped: dict = {}
    for target_key, source_key in field_mapping.items():
        if source_key and source_key in raw_data:
            mapped[target_key] = raw_data[source_key]
        else:
            mapped[target_key] = None
    return mapped


def _filter_by_sheets(raw: dict, selected_sheets: Any) -> bool:
    if not selected_sheets or not isinstance(selected_sheets, (list, tuple, set)):
        return True
    return raw.get("sheet_name") in selected_sheets


def import_review(request: ImportConfirmRequest, current_user: dict) -> dict:
    """Validate batch and return conflicts without creating customers."""
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        batch = session.fetch_one("SELECT * FROM customer_source_batches WHERE id = ?", (request.batch_id,))
        if not batch:
            raise ValueError("batch_not_found")
        if batch["status"] not in {"preview", "partial"}:
            raise ValueError("batch_not_confirmed")

        raw_rows = session.fetch_all(
            "SELECT * FROM customer_raw_records WHERE batch_id = ?",
            (request.batch_id,),
        )
        if not raw_rows:
            raise ValueError("batch_empty")

        errors = []
        conflicts = []
        for raw in raw_rows:
            if not _filter_by_sheets(raw, request.selected_sheets):
                continue
            raw_data = parse_json(raw["raw_data"], {})
            raw_data = _apply_field_mapping(raw_data, request.field_mapping)
            try:
                _normalize_customer_from_row(raw_data, request.batch_id)
            except ValueError as exc:
                errors.append({"row": raw["row_number"], "sheet": raw["sheet_name"], "error": str(exc)})
                continue

            candidates = _detect_duplicate_candidates(session, raw_data, request.batch_id)
            if candidates:
                conflicts.append({
                    "raw_record_id": raw["id"],
                    "row_number": raw["row_number"],
                    "sheet_name": raw["sheet_name"],
                    "candidates": candidates,
                    "resolution": None,
                })

        return {
            "batch_id": request.batch_id,
            "errors": errors,
            "conflicts": conflicts,
            "can_confirm": not errors and not conflicts,
        }
    finally:
        conn.close()


def import_confirm(request: ImportConfirmRequest, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        batch = session.fetch_one("SELECT * FROM customer_source_batches WHERE id = ?", (request.batch_id,))
        if not batch:
            raise ValueError("batch_not_found")
        if batch["status"] not in {"preview", "partial"}:
            raise ValueError("batch_not_confirmed")

        raw_rows = session.fetch_all(
            "SELECT * FROM customer_raw_records WHERE batch_id = ?",
            (request.batch_id,),
        )
        if not raw_rows:
            raise ValueError("batch_empty")

        imported = 0
        skipped = 0
        errors: list[dict] = []

        with session.transaction():
            for raw in raw_rows:
                if not _filter_by_sheets(raw, request.selected_sheets):
                    continue
                raw_data = parse_json(raw["raw_data"], {})
                raw_data = _apply_field_mapping(raw_data, request.field_mapping)
                try:
                    normalized = _normalize_customer_from_row(raw_data, request.batch_id)
                except ValueError as exc:
                    errors.append({"row": raw["row_number"], "sheet": raw["sheet_name"], "error": str(exc)})
                    session.update(
                        "customer_raw_records",
                        raw["id"],
                        {
                            "validation_errors": json.dumps([str(exc)], ensure_ascii=False),
                            "conflict_resolution": "error",
                        },
                    )
                    continue

                candidates = _detect_duplicate_candidates(session, raw_data, request.batch_id)
                customer_id = None
                conflict_resolution = "created"

                if candidates:
                    high_confidence = [c for c in candidates if c["match_type"] in {"tax_id_or_license", "email", "phone", "website"}]
                    if high_confidence and request.duplicate_policy == "merge":
                        customer_id = high_confidence[0]["customer"]["id"]
                        conflict_resolution = "merged"
                        updates = {k: v for k, v in normalized.items() if k not in {"raw_data", "status", "data_status", "verification_status", "crm_status", "activity_status"}}
                        session.update("customers", customer_id, updates)
                    elif high_confidence and request.duplicate_policy == "skip":
                        skipped += 1
                        conflict_resolution = "skipped"
                    elif high_confidence and request.duplicate_policy == "error":
                        errors.append({"row": raw["row_number"], "sheet": raw["sheet_name"], "error": "duplicate_detected"})
                        conflict_resolution = "error"
                    else:
                        if request.duplicate_policy == "skip":
                            skipped += 1
                            conflict_resolution = "skipped"
                        elif request.duplicate_policy == "create_new":
                            customer_id = _create_customer_from_normalized(session, normalized, current_user)
                            conflict_resolution = "created"
                        elif request.duplicate_policy == "merge" and candidates:
                            customer_id = candidates[0]["customer"]["id"]
                            conflict_resolution = "merged"
                            updates = {k: v for k, v in normalized.items() if k not in {"raw_data", "status", "data_status", "verification_status", "crm_status", "activity_status"}}
                            session.update("customers", customer_id, updates)
                        elif request.duplicate_policy == "error":
                            errors.append({"row": raw["row_number"], "sheet": raw["sheet_name"], "error": "duplicate_detected"})
                            conflict_resolution = "error"
                        else:
                            customer_id = _create_customer_from_normalized(session, normalized, current_user)
                            conflict_resolution = "created"
                else:
                    customer_id = _create_customer_from_normalized(session, normalized, current_user)
                    conflict_resolution = "created"

                if customer_id:
                    imported += 1
                    session.update(
                        "customer_raw_records",
                        raw["id"],
                        {
                            "normalized_customer_id": customer_id,
                            "conflict_resolution": conflict_resolution,
                            "conflict_details": json.dumps({"candidates": candidates}, ensure_ascii=False),
                        },
                    )

            new_status = "failed" if errors else "confirmed"
            session.update(
                "customer_source_batches",
                request.batch_id,
                {
                    "status": new_status,
                    "success_count": imported,
                    "skip_count": skipped,
                    "error_count": len(errors),
                },
            )
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="import_confirm", entity_type="customer_batch", entity_id=request.batch_id, details=f"imported={imported} skipped={skipped} errors={len(errors)}"),
        )
        return {"batch_id": request.batch_id, "imported": imported, "skipped": skipped, "errors": errors}
    finally:
        conn.close()


def import_cancel(batch_id: int, current_user: dict) -> dict:
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        batch = session.fetch_one("SELECT * FROM customer_source_batches WHERE id = ?", (batch_id,))
        if not batch:
            raise ValueError("batch_not_found")
        if batch["status"] not in {"preview", "partial"}:
            return {"message": "No changes"}
        with session.transaction():
            session.update("customer_source_batches", batch_id, {"status": "cancelled"})
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="import_cancel", entity_type="customer_batch", entity_id=batch_id),
        )
        return {"message": "Import cancelled successfully"}
    finally:
        conn.close()


def import_cleanup(current_user: dict) -> dict:
    """Mark expired preview batches as expired. Idempotent."""
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        now = now_iso()
        with session.transaction():
            expired_batches = session.fetch_all(
                "SELECT id FROM customer_source_batches WHERE status = 'preview' AND expires_at IS NOT NULL AND expires_at <= ?",
                (now,),
            )
            for batch in expired_batches:
                session.update("customer_source_batches", batch["id"], {"status": "expired"})
        log_audit(
            current_user=current_user,
            data=AuditLogCreate(action="import_cleanup", entity_type="customer_batch", entity_id=None, details=f"expired={len(expired_batches)}"),
        )
        return {"expired": len(expired_batches)}
    finally:
        conn.close()


def import_customers(file: io.BytesIO, filename: str, current_user: dict) -> dict:
    if not filename.endswith('.csv'):
        raise ValueError("Only CSV files are allowed")
    content = file.read()
    rows, _ = _parse_csv(content)
    conn = get_db()
    try:
        session = DatabaseSession(conn)
        now = now_iso()
        imported = 0
        with session.transaction():
            for row in rows:
                try:
                    normalized = _normalize_customer_from_row(row, 0)
                except ValueError:
                    continue
                customer_id = _create_customer_from_normalized(session, normalized, current_user)
                imported += 1
        return {"message": f"Imported {imported} customers successfully", "count": imported}
    finally:
        conn.close()
